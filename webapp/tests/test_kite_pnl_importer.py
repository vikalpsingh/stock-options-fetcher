from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook

from kite_pnl_importer import KitePnlImportRepository, parse_kite_pnl_xlsx
from monthly_risk_control import MonthlyRiskRepository


def _workbook_bytes(*, account="UM4139", start="2026-04-01", end="2026-10-17") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "F&O"
    sheet["B7"] = "Client ID"
    sheet["C7"] = account
    sheet["B11"] = f"P&L Statement for F&O from {start} to {end}"
    sheet["B15"] = "Charges"
    sheet["C15"] = 150.0
    sheet["B16"] = "Other Credit & Debit"
    sheet["C16"] = 10.0
    sheet["B17"] = "Realized P&L"
    sheet["C17"] = 1500.0
    sheet["B18"] = "Unrealized P&L"
    sheet["C18"] = -50.0
    headers = ["Symbol", "ISIN", "Quantity", "Buy Value", "Sell Value", "Realized P&L", "Realized P&L Pct.", "Previous Closing Price", "Open Quantity", "Open Quantity Type", "Open Value", "Unrealized P&L", "Unrealized P&L Pct."]
    for column, value in enumerate(headers, 2):
        sheet.cell(38, column, value)
    sheet.append([])
    rows = [
        ["ABC26MAY100CE", "", 100, 1000, 1800, 800, 80, 0, 0, "", 0, 0, 0],
        ["XYZ26JUN200PE", "", 100, 2000, 2700, 700, 35, 0, 0, "", 0, 0, 0],
        ["OPEN26NOV300CE", "", 0, 0, 0, 0, 0, 10, 100, "buy", 1000, -50, -5],
    ]
    for row_index, values in enumerate(rows, 39):
        for column, value in enumerate(values, 2):
            sheet.cell(row_index, column, value)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_parse_kite_pnl_reconciles_and_keeps_open_pnl_in_asof_month():
    parsed = parse_kite_pnl_xlsx(_workbook_bytes())

    assert parsed.account_id == "UM4139"
    assert parsed.gross_realized_pnl == 1500
    assert parsed.charges == 150
    assert round(sum(row["net_realized_pnl"] for row in parsed.monthly_rows), 2) == 1360
    assert [row["month_key"] for row in parsed.monthly_rows] == ["2026-05", "2026-06", "2026-10"]
    assert sum(row["unrealized_pnl"] for row in parsed.monthly_rows) == -50
    assert parsed.allocation_method == "CONTRACT_MONTH_INFERRED"


def test_kite_import_is_idempotent_and_builds_quarters(tmp_path):
    parsed = parse_kite_pnl_xlsx(_workbook_bytes())
    repository = KitePnlImportRepository(tmp_path / "pnl.sqlite3")

    first = repository.import_workbook(parsed, "pnl.xlsx")
    duplicate = repository.import_workbook(parsed, "pnl.xlsx")
    quarters = {row["quarter"]: row for row in repository.quarterly_rows("UM4139")}

    assert first["duplicate"] is False
    assert duplicate["duplicate"] is True
    assert len(repository.monthly_rows("UM4139")) == 3
    assert round(quarters["2026-Q2"]["net_realized_pnl"], 2) == 1360


def test_imported_month_is_authoritative_for_monthly_risk(tmp_path):
    db_path = tmp_path / "pnl.sqlite3"
    parsed = parse_kite_pnl_xlsx(_workbook_bytes())
    KitePnlImportRepository(db_path).import_workbook(parsed, "pnl.xlsx")

    snapshot = MonthlyRiskRepository(db_path).latest_snapshot("UM4139", date(2026, 5, 31))

    may = next(row for row in parsed.monthly_rows if row["month_key"] == "2026-05")
    assert snapshot["booked_pnl"] == may["net_realized_pnl"]
    assert "KITE_PNL_IMPORT" in snapshot["data_status"]
