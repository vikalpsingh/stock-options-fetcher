"""Parse Kite P&L XLSX exports into duplicate-safe monthly account totals."""

from __future__ import annotations

import hashlib
import io
import re
import sqlite3
from calendar import monthrange
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - deployment dependency diagnostic
    load_workbook = None


MONTHS = {name: index for index, name in enumerate(
    ("JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"), 1
)}


def _number(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _month_key_from_symbol(symbol: str, fallback: str) -> str:
    text = str(symbol or "").upper().replace(" ", "")
    monthly = re.search(r"(\d{2})(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC)", text)
    if monthly:
        return f"20{monthly.group(1)}-{MONTHS[monthly.group(2)]:02d}"
    weekly = re.search(r"(?:NIFTY|BANKNIFTY|FINNIFTY)(\d{2})(\d)(\d{2})", text)
    if weekly:
        month = int(weekly.group(2))
        if 1 <= month <= 9:
            return f"20{weekly.group(1)}-{month:02d}"
    return fallback


@dataclass(frozen=True)
class KitePnlWorkbook:
    account_id: str
    period_start: date
    period_end: date
    gross_realized_pnl: float
    charges: float
    other_credit_debit: float
    unrealized_pnl: float
    monthly_rows: tuple[dict[str, Any], ...]
    allocation_method: str
    source_hash: str


def parse_kite_pnl_xlsx(content: bytes) -> KitePnlWorkbook:
    if not content:
        raise ValueError("The Kite P&L workbook is empty.")
    if load_workbook is None:
        raise RuntimeError("Kite P&L import requires openpyxl. Install requirements.txt on the server.")
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Could not read the Kite P&L .xlsx file: {exc}") from exc
    sheet = workbook["F&O"] if "F&O" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    sheet.reset_dimensions()
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    account_id = ""
    period_start: date | None = None
    period_end: date | None = None
    summary: dict[str, float] = {}
    header_index = -1
    for index, row in enumerate(rows):
        for col, value in enumerate(row):
            label = str(value or "").strip()
            if label == "Client ID" and col + 1 < len(row):
                account_id = str(row[col + 1] or "").strip().upper()
            period_match = re.search(r"from\s+(\d{4}-\d{2}-\d{2})\s+to\s+(\d{4}-\d{2}-\d{2})", label, re.I)
            if period_match:
                period_start = date.fromisoformat(period_match.group(1))
                period_end = date.fromisoformat(period_match.group(2))
            if label in {"Charges", "Other Credit & Debit", "Realized P&L", "Unrealized P&L"} and col + 1 < len(row):
                value_number = row[col + 1]
                if isinstance(value_number, (int, float)):
                    summary[label] = _number(value_number)
        labels = [str(value or "").strip() for value in row]
        if "Symbol" in labels and "Realized P&L" in labels and "Buy Value" in labels:
            header_index = index
            header = labels
            break
    if not account_id or not period_start or not period_end or header_index < 0:
        raise ValueError("This is not a supported Kite F&O P&L workbook: account, period, or Symbol table is missing.")
    indexes = {name: header.index(name) for name in ("Symbol", "Buy Value", "Sell Value", "Realized P&L", "Unrealized P&L")}
    fallback_month = period_end.strftime("%Y-%m")
    by_month: dict[str, dict[str, float]] = {}
    for row in rows[header_index + 1:]:
        symbol = str(row[indexes["Symbol"]] or "").strip().upper() if len(row) > indexes["Symbol"] else ""
        if not symbol:
            continue
        realised = _number(row[indexes["Realized P&L"]])
        buy_value = _number(row[indexes["Buy Value"]])
        sell_value = _number(row[indexes["Sell Value"]])
        open_pnl = _number(row[indexes["Unrealized P&L"]])
        # Open P&L belongs to the statement as-of month, not the future contract
        # expiry month. Closed/turnover rows use contract month when dates are absent.
        key = _month_key_from_symbol(symbol, fallback_month) if (realised or buy_value or sell_value) else fallback_month
        bucket = by_month.setdefault(key, {"gross_realized_pnl": 0.0, "unrealized_pnl": 0.0, "turnover": 0.0, "row_count": 0.0})
        bucket["gross_realized_pnl"] += realised
        bucket["unrealized_pnl"] += open_pnl
        bucket["turnover"] += abs(buy_value) + abs(sell_value)
        bucket["row_count"] += 1
    if not by_month:
        raise ValueError("No F&O P&L rows were found in the Kite workbook.")
    gross_reported = round(summary.get("Realized P&L", sum(v["gross_realized_pnl"] for v in by_month.values())), 2)
    row_total = sum(v["gross_realized_pnl"] for v in by_month.values())
    if abs(row_total - gross_reported) > 1.0:
        raise ValueError(f"Kite workbook does not reconcile: summary realised P&L {gross_reported:.2f}, rows {row_total:.2f}.")
    charges = round(summary.get("Charges", 0.0), 2)
    other = round(summary.get("Other Credit & Debit", 0.0), 2)
    total_turnover = sum(v["turnover"] for v in by_month.values())
    keys = sorted(by_month)
    monthly_rows: list[dict[str, Any]] = []
    allocated_charges = 0.0
    allocated_other = 0.0
    single_month = period_start.strftime("%Y-%m") == period_end.strftime("%Y-%m")
    method = "STATEMENT_MONTH_EXACT" if single_month else "CONTRACT_MONTH_INFERRED"
    if single_month and len(keys) != 1:
        combined = {
            "gross_realized_pnl": sum(v["gross_realized_pnl"] for v in by_month.values()),
            "unrealized_pnl": sum(v["unrealized_pnl"] for v in by_month.values()),
            "turnover": total_turnover,
            "row_count": sum(v["row_count"] for v in by_month.values()),
        }
        by_month = {period_end.strftime("%Y-%m"): combined}
        keys = list(by_month)
    for position, key in enumerate(keys):
        bucket = by_month[key]
        if position == len(keys) - 1:
            charge_share = round(charges - allocated_charges, 2)
            other_share = round(other - allocated_other, 2)
        else:
            weight = (bucket["turnover"] / total_turnover) if total_turnover else (1 / len(keys))
            charge_share = round(charges * weight, 2)
            other_share = round(other * weight, 2)
            allocated_charges += charge_share
            allocated_other += other_share
        gross = round(bucket["gross_realized_pnl"], 2)
        monthly_rows.append({
            "month_key": key,
            "gross_realized_pnl": gross,
            "charges": charge_share,
            "other_credit_debit": other_share,
            "net_realized_pnl": round(gross - charge_share + other_share, 2),
            "unrealized_pnl": round(bucket["unrealized_pnl"], 2),
            "turnover": round(bucket["turnover"], 2),
            "row_count": int(bucket["row_count"]),
            "allocation_method": method,
        })
    return KitePnlWorkbook(
        account_id=account_id,
        period_start=period_start,
        period_end=period_end,
        gross_realized_pnl=gross_reported,
        charges=charges,
        other_credit_debit=other,
        unrealized_pnl=round(summary.get("Unrealized P&L", 0.0), 2),
        monthly_rows=tuple(monthly_rows),
        allocation_method=method,
        source_hash=hashlib.sha256(content).hexdigest(),
    )


class KitePnlImportRepository:
    def __init__(self, db_path: str | Path):
        self.db_path = Path(db_path)

    def connect(self) -> sqlite3.Connection:
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        return conn

    def ensure_schema(self) -> None:
        with self.connect() as conn:
            conn.executescript("""
                CREATE TABLE IF NOT EXISTS pnl_kite_import (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    account_id TEXT NOT NULL, broker_client_id TEXT NOT NULL DEFAULT '', file_name TEXT NOT NULL,
                    source_hash TEXT NOT NULL UNIQUE, period_start TEXT NOT NULL,
                    period_end TEXT NOT NULL, gross_realized_pnl REAL NOT NULL,
                    charges REAL NOT NULL, other_credit_debit REAL NOT NULL,
                    unrealized_pnl REAL NOT NULL, allocation_method TEXT NOT NULL,
                    imported_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS pnl_kite_monthly (
                    account_id TEXT NOT NULL, month_key TEXT NOT NULL,
                    import_id INTEGER NOT NULL, gross_realized_pnl REAL NOT NULL,
                    charges REAL NOT NULL, other_credit_debit REAL NOT NULL,
                    net_realized_pnl REAL NOT NULL, unrealized_pnl REAL NOT NULL,
                    turnover REAL NOT NULL, row_count INTEGER NOT NULL,
                    allocation_method TEXT NOT NULL, updated_at TEXT NOT NULL,
                    PRIMARY KEY(account_id, month_key)
                );
            """)
            columns = {str(row[1]) for row in conn.execute("PRAGMA table_info(pnl_kite_import)").fetchall()}
            if "broker_client_id" not in columns:
                conn.execute("ALTER TABLE pnl_kite_import ADD COLUMN broker_client_id TEXT NOT NULL DEFAULT ''")

    def import_workbook(self, parsed: KitePnlWorkbook, file_name: str, account_id: str | None = None) -> dict[str, Any]:
        self.ensure_schema()
        account = str(account_id or parsed.account_id).strip().upper()
        if not account:
            raise ValueError("Select an account for this Kite P&L import.")
        now = datetime.now().astimezone().isoformat(timespec="seconds")
        with self.connect() as conn:
            existing = conn.execute("SELECT id FROM pnl_kite_import WHERE source_hash=?", (parsed.source_hash,)).fetchone()
            if existing:
                return {"duplicate": True, "import_id": int(existing["id"]), "account_id": account, "months": len(parsed.monthly_rows)}
            cursor = conn.execute(
                """INSERT INTO pnl_kite_import(account_id,broker_client_id,file_name,source_hash,period_start,period_end,
                   gross_realized_pnl,charges,other_credit_debit,unrealized_pnl,allocation_method,imported_at)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
                (account, parsed.account_id, Path(file_name).name, parsed.source_hash, parsed.period_start.isoformat(), parsed.period_end.isoformat(),
                 parsed.gross_realized_pnl, parsed.charges, parsed.other_credit_debit, parsed.unrealized_pnl,
                 parsed.allocation_method, now),
            )
            import_id = int(cursor.lastrowid)
            for row in parsed.monthly_rows:
                conn.execute(
                    """INSERT INTO pnl_kite_monthly VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
                       ON CONFLICT(account_id,month_key) DO UPDATE SET
                       import_id=excluded.import_id,gross_realized_pnl=excluded.gross_realized_pnl,
                       charges=excluded.charges,other_credit_debit=excluded.other_credit_debit,
                       net_realized_pnl=excluded.net_realized_pnl,unrealized_pnl=excluded.unrealized_pnl,
                       turnover=excluded.turnover,row_count=excluded.row_count,
                       allocation_method=excluded.allocation_method,updated_at=excluded.updated_at""",
                    (account, row["month_key"], import_id, row["gross_realized_pnl"], row["charges"],
                     row["other_credit_debit"], row["net_realized_pnl"], row["unrealized_pnl"],
                     row["turnover"], row["row_count"], row["allocation_method"], now),
                )
        return {"duplicate": False, "import_id": import_id, "account_id": account, "months": len(parsed.monthly_rows)}

    def monthly_rows(self, account_id: str) -> list[dict[str, Any]]:
        self.ensure_schema()
        with self.connect() as conn:
            return [dict(row) for row in conn.execute(
                "SELECT * FROM pnl_kite_monthly WHERE account_id=? ORDER BY month_key DESC", (account_id,)
            ).fetchall()]

    def account_ids(self) -> list[str]:
        self.ensure_schema()
        with self.connect() as conn:
            return [str(row["account_id"]) for row in conn.execute(
                "SELECT DISTINCT account_id FROM pnl_kite_monthly ORDER BY account_id"
            ).fetchall()]

    def quarterly_rows(self, account_id: str) -> list[dict[str, Any]]:
        rows = self.monthly_rows(account_id)
        grouped: dict[str, dict[str, Any]] = {}
        for row in rows:
            year, month = (int(part) for part in str(row["month_key"]).split("-"))
            key = f"{year}-Q{((month - 1) // 3) + 1}"
            item = grouped.setdefault(key, {"quarter": key, "gross_realized_pnl": 0.0, "charges": 0.0, "net_realized_pnl": 0.0, "months": 0})
            for field in ("gross_realized_pnl", "charges", "net_realized_pnl"):
                item[field] += _number(row[field])
            item["months"] += 1
        for item in grouped.values():
            for field in ("gross_realized_pnl", "charges", "net_realized_pnl"):
                item[field] = round(item[field], 2)
        return [grouped[key] for key in sorted(grouped, reverse=True)]
